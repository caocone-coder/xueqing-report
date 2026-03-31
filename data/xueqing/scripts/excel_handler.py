#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 读写模块
处理源文件读取和目标文件写入
"""

import openpyxl
from datetime import datetime
import shutil


class ExcelHandler:
    def __init__(self, source_file, target_file):
        self.source_file = source_file
        self.target_file = target_file

    def read_source_urls(self, column='K', skip_header=True):
        """
        读取源文件指定列的 URL

        Args:
            column: 列名（如 'K'）
            skip_header: 是否跳过标题行

        Returns:
            dict: {row_num: url} 字典
        """
        wb = openpyxl.load_workbook(self.source_file)
        ws = wb.active

        # 将列名转换为列索引
        col_idx = openpyxl.utils.column_index_from_string(column)

        url_dict = {}
        start_row = 2 if skip_header else 1

        for row in range(start_row, ws.max_row + 1):
            cell = ws.cell(row, col_idx)
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('http'):
                url_dict[row] = cell.value

        wb.close()
        print(f"从源文件读取到 {len(url_dict)} 个 URL")
        return url_dict

    def read_target_headers(self):
        """
        读取目标文件的标题行

        Returns:
            list: 字段名列表
        """
        wb = openpyxl.load_workbook(self.target_file)
        ws = wb.active

        headers = []
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(1, col)
            if cell.value:
                headers.append(cell.value)

        wb.close()
        print(f"目标文件包含 {len(headers)} 个字段")
        return headers

    def get_target_next_row(self):
        """
        获取目标文件下一个可写入的行号

        Returns:
            int: 下一行行号
        """
        wb = openpyxl.load_workbook(self.target_file)
        ws = wb.active
        next_row = ws.max_row + 1
        wb.close()
        return next_row

    def append_data_to_target(self, data_list):
        """
        追加数据到目标文件

        Args:
            data_list: 数据列表，每个元素是一个包含 18 个字段的列表

        Returns:
            bool: 是否成功
        """
        try:
            wb = openpyxl.load_workbook(self.target_file)
            ws = wb.active

            initial_row_count = ws.max_row

            for data_row in data_list:
                ws.append(data_row)

            wb.save(self.target_file)
            wb.close()

            final_row_count = initial_row_count + len(data_list)
            print(f"成功写入 {len(data_list)} 行数据")
            print(f"目标文件行数: {initial_row_count} -> {final_row_count}")
            return True

        except Exception as e:
            print(f"写入失败: {e}")
            return False

    def create_backup(self, backup_dir):
        """
        创建目标文件备份

        Args:
            backup_dir: 备份目录

        Returns:
            str: 备份文件路径
        """
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_name = f"model 配置表_backup_{timestamp}.xlsx"
        backup_path = f"{backup_dir}/{backup_name}"

        shutil.copy2(self.target_file, backup_path)
        print(f"备份已创建: {backup_path}")
        return backup_path
